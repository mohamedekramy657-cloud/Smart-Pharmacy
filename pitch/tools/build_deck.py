"""Builds the Smart Pharmacy hackathon pitch deck (16:9) with native, editable charts.

Data: cited public sources (see SOURCES), team measurements, and live demo analytics queried
from the prototype database (clearly labelled as demo data). Screens: real prototype captures.
"""
import json
import os
import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
MOCK = ROOT / "assets" / "mockups"
OUT = ROOT / "Smart_Pharmacy_Hackathon_Pitch.pptx"

NAVY, NAVY2, TEAL, TEAL_D, TEAL_L, MINT = "0B1F2A", "12303F", "0F766E", "115E59", "2DD4BF", "CCFBF1"
AMBER, AMBER_L, CORAL, CORAL_L, SKY = "F59E0B", "FEF3C7", "E11D48", "FFE4E6", "0284C7"
SLATE, MUTED, BORDER, LIGHT, WHITE = "0F172A", "64748B", "E2E8F0", "F8FAFC", "FFFFFF"
FONT, AR_FONT = "Calibri", "Arial"

# ---------------------------------------------------------------- data
SOURCES = [
    ("CAPMAS population clock via EgyptToday (19 Sep 2026): 109.5M residents; Ismailia 1.5M", "https://www.egypttoday.com/Article/1/149579/Egypt%E2%80%99s-domestic-population-reaches-109-5-Million-CAPMAS"),
    ("Egyptian Pharmacists Syndicate records, in Expl. Res. Clin. Soc. Pharmacy (2023): 95,000 pharmacies; 313k pharmacists", "https://www.sciencedirect.com/science/article/pii/S2667276623001658"),
    ("Egyptian Drug Authority via Daily News Egypt (3 Jun 2025): market EGP 216bn (2023) → 309bn (2024)", "https://www.dailynewsegypt.com/2025/06/03/egypts-pharmaceutical-market-hits-egp-309bn-in-2024-marking-43-growth/"),
    ("Egyptian Drug Authority via Arab Finance (Jul 2026): EGP 422bn in 2025 (+37%), ~4bn packs", "https://www.arabfinance.com/en/news/newdetails/exploring-egypt-rising-pharmaceutical-industry"),
    ("Middle East Monitor (30 Jul 2024): ~1,000 of 17,000 registered drugs short; Concor, Brufen, Cataflam; 30M+ hypertensive", "https://www.middleeastmonitor.com/20240730-egyptians-face-illness-amid-a-serious-scarcity-of-medicines/"),
    ("Ahram Online (20 Aug 2024): shortages fell from ~1,000 to ~500 medicine types", "https://english.ahram.org.eg/NewsContentP/1/530363/Egypt/Egypt-drug-shortage-crisis-Major-relief-around-the.aspx"),
    ("Business Monthly (23 Oct 2024): 800+ types short (FEI memo); 20–40% of trade names (CICMS)", "https://businessmonthlyeg.com/pharmaceutical-shortages-worsen-in-egypt-how-government-and-pharmacists-respond/"),
    ("Zawia3 (25 Apr 2025): 1,471 branded drugs unavailable, 300 without alternatives (FEI Pharma Chamber, Mar 2025)", "https://zawia3.com/en/psychiatric-medications/"),
    ("EnterpriseAM (3 Sep 2025): ~200 missing medicines in 2025; chronic & life-saving gaps persist", "https://enterpriseam.com/egypt/2025/09/03/pharma-firms-push-for-price-hikes-as-shortages-and-black-market-resurface/"),
    ("Al Majalla (9 Dec 2025): shortages cut from 40% to under 5% of total supply", "https://en.majalla.com/node/328622/business-economy/egypt%E2%80%99s-drug-price-increases-help-companies-not-patients"),
    ("The National (15 Feb 2024): Augmentin & chronic-drug shortages; ~11M with diabetes; black-market risk", "https://www.thenationalnews.com/mena/egypt/2024/02/15/medicine-shortages-cairo-diabetes-drugs/"),
    ("Wamda (19 Mar 2024): ~30% of pharmacists' time on manual stock checks; 7,000+ products, 3,000 suppliers", "https://www.wamda.com/2024/03/egypt-pharmacy-marts-secures-bridge-round-acasia-ventures"),
    ("Mada Masr (9 Mar 2021): law allows a pharmacist to own no more than two pharmacies", "https://www.madamasr.com/en/2021/03/09/feature/economy/the-rise-and-fall-of-egypts-top-pharmaceutical-chain/"),
    ("Naos Solutions – Egyptians & Digital 2025 (DataReportal data): 96.3M internet users, 81.9%", "https://naos-solutions.com/egyptians-and-digital-2025-report/"),
    ("IDSC via Zawya (20 Apr 2026): 92.6M mobile-internet subscriptions (Dec 2025)", "https://www.zawya.com/en/business/technology-and-telecom/egypts-mobile-internet-users-hit-926mln-in-december-2025-idsc-hwrcfjnz"),
    ("Ahram Online: Egyptian Drug Authority 'Availability' service for reporting missing medicines", "https://english.ahram.org.eg/UI/Front/Inner.aspx?NewsContentID=527162"),
    ("Competitor profiles – AccretiveEdge (Aug 2025: Yodawy, Chefaa, Vezeeta); Tracxn (Pharmacy Marts, iSupply)", "https://accretiveedge.com/articles/healthtech-startups-middle-east/"),
    ("Smart Pharmacy prototype – team measurements, Oct 2026 (tests, latency on synthetic data, demo analytics)", None),
]


def demo_analytics():
    """Live analytics from the prototype database (demo data); cached for offline rebuilds."""
    cache = ROOT / "data" / "demo_analytics.json"
    sql = {
        "outcomes": "SELECT outcome, count(*) FROM search_events GROUP BY 1",
        "bands": "SELECT coalesce(matched_radius_km::int::text,'none'), count(*) FROM search_events GROUP BY 1",
        "freshness": """SELECT CASE WHEN i.last_update_source='INTEGRATION' THEN 'integration' WHEN i.last_updated_at > now()-interval '12 hours' THEN 'confirmed'
                         WHEN i.last_updated_at > now()-interval '72 hours' THEN 'recent' ELSE 'stale' END, count(*)
                         FROM inventory_items i JOIN organizations o ON o.id=i.organization_id WHERE o.type='PHARMACY' GROUP BY 1""",
        "unmet": """SELECT m.brand_name_en || ' ' || m.strength, count(*) FROM search_events e JOIN medicines m ON m.id=e.medicine_id
                     WHERE e.outcome <> 'FOUND' GROUP BY 1 ORDER BY 2 DESC LIMIT 3""",
        "pharmacies": "SELECT 'n', count(*) FROM organizations WHERE type='PHARMACY'",
    }
    try:
        env = dict(os.environ, PGPASSWORD="pharma_dev_password")
        data = {}
        for k, q in sql.items():
            out = subprocess.run(["psql", "-h", "127.0.0.1", "-U", "pharma", "-d", "pharma_dev", "-tA", "-F", "|", "-c", q], capture_output=True, text=True, env=env, check=True).stdout
            data[k] = {line.split("|")[0]: int(line.split("|")[1]) for line in out.strip().splitlines() if "|" in line}
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(json.dumps(data, indent=2))
        return data
    except Exception:
        return json.loads(cache.read_text())


# ---------------------------------------------------------------- helpers
def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, size, color, bold=False, italic=False, font=FONT, spc=None):
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    for tag, face in (("a:cs", AR_FONT), ("a:ea", font)):
        if rPr.find(qn(tag)) is None:
            el = etree.Element(qn(tag))
            el.set("typeface", face)
            latin.addnext(el)
    if spc:
        rPr.set("spc", str(spc))


def text(slide, x, y, w, h, content, size=14, color=SLATE, bold=False, italic=False, align="l", anchor="t", spc=None, line=1.0, after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line
        p.space_after = Pt(after)
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = t
            style_run(run, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", italic), o.get("font", FONT), o.get("spc", spc))
            if o.get("link"):
                run.hyperlink.address = o["link"]
    return tb


def set_alpha(shape, alpha_pct):
    clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))


def add_shadow(shape, blur=14, dist=3, alpha=16):
    shape.shadow.inherit = False
    eff = shape._element.spPr.find(qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur * 12700), dist=str(dist * 12700), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="0F172A")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def box(slide, x, y, w, h, fill=None, line=None, radius=0.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.0, alpha=None, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            set_alpha(s, alpha)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    if shadow:
        add_shadow(s)
    else:
        s.shadow.inherit = False
    return s


def label(shape, t, size=12, color=WHITE, bold=True, align="c"):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align]
    run = p.add_run()
    run.text = t
    style_run(run, size, color, bold)


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = c.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def pic(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(MOCK / name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def header(slide, kicker, title, dark=False):
    text(slide, 0.6, 0.42, 10, 0.3, kicker.upper(), size=11, color=TEAL_L if dark else TEAL, bold=True, spc=200)
    text(slide, 0.6, 0.68, 12.2, 0.8, title, size=28, bold=True, color=WHITE if dark else NAVY)


def footer(slide, n, sources=None, dark=False):
    c = "94A3B8" if dark else MUTED
    text(slide, 0.6, 7.08, 6, 0.25, [[("Smart Pharmacy · ", {}), ("الصيدلية الذكية", {})]], size=9, color=c)
    text(slide, 11.9, 7.08, 0.85, 0.25, str(n), size=9, color=c, align="r")
    if sources:
        text(slide, 3.6, 7.08, 8.0, 0.25, "Sources: " + sources, size=9, color=c, italic=True, align="r")


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def kpi(slide, x, y, w, h, value, desc, color=TEAL, fill=WHITE, vsize=30, dsize=12, src=None):
    box(slide, x, y, w, h, fill=fill, line=BORDER, radius=0.08, shadow=True)
    text(slide, x + 0.22, y + 0.16, w - 0.44, 0.6, value, size=vsize, bold=True, color=color)
    text(slide, x + 0.22, y + 0.2 + vsize / 72 * 1.12, w - 0.44, h - 0.3 - vsize / 72, desc, size=dsize, color=SLATE, line=1.05)
    if src:
        text(slide, x + 0.22, y + h - 0.3, w - 0.44, 0.22, src, size=8.5, color=MUTED, italic=True)


def chart_base(ch, size=11):
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(size)
    ch.font.color.rgb = rgb(SLATE)


def column_chart(slide, x, y, w, h, cats, vals, colors, fmt="#,##0", lsize=13, csize=11, gap=55, horizontal=False):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Series", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    chart_base(ch, csize)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.font.size = Pt(lsize)
    dl.font.bold = True
    dl.font.color.rgb = rgb(NAVY)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = False
    va.minimum_scale = 0
    va.maximum_scale = max(vals) * (1.25 if horizontal else 1.18) if max(vals) else 1
    va.visible = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(csize)
    ca.format.line.color.rgb = rgb(BORDER)
    ca.has_major_gridlines = False
    if horizontal:
        ca.reverse_order = True
    return ch


def doughnut(slide, x, y, w, h, cats, vals, colors, hole=58):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Series", vals)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart_base(ch, 11)
    plot = ch.plots[0]
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_percentage = True
    dl.show_value = False
    dl.show_category_name = False
    dl.number_format = "0%"
    dl.number_format_is_linked = False
    dl.font.size = Pt(12)
    dl.font.bold = True
    dl.font.color.rgb = rgb(WHITE)
    hs = plot._element.find(qn("c:holeSize"))
    if hs is None:
        hs = etree.SubElement(plot._element, qn("c:holeSize"))
    hs.set("val", str(hole))
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    return ch


def grouped_chart(slide, x, y, w, h, cats, series, colors, fmt="0"):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart_base(ch, 11)
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.overlap = -10
    for s, c in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.font.size = Pt(12)
    dl.font.bold = True
    dl.font.color.rgb = rgb(NAVY)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = False
    va.minimum_scale = 0
    va.maximum_scale = max(max(v) for _, v in series) * 1.2
    va.visible = False
    ch.category_axis.format.line.color.rgb = rgb(BORDER)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    return ch


def stacked_bar(slide, x, y, w, h, series, colors):
    cd = CategoryChartData()
    cd.categories = ["Tests"]
    for name, v in series:
        cd.add_series(name, [v])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart_base(ch, 10)
    plot = ch.plots[0]
    plot.gap_width = 30
    for s, c in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(12)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = rgb(WHITE)
    plot.data_labels.position = XL_LABEL_POSITION.CENTER
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.visible = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    return ch


def bullet_list(slide, x, y, w, items, size=13, color=SLATE, dot=TEAL, gap=0.42, dot_size=0.12):
    for i, it in enumerate(items):
        yy = y + i * gap
        box(slide, x, yy + 0.07, dot_size, dot_size, fill=dot, shape=MSO_SHAPE.OVAL)
        text(slide, x + 0.25, yy, w - 0.25, gap, it, size=size, color=color, line=1.0)


# ---------------------------------------------------------------- slides
def build():
    A = demo_analytics()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    n = 0

    def new(dark=False, bg=None):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg or (NAVY if dark else WHITE))
        return s

    # 1 — Cover
    s = new(dark=True)
    box(s, 8.6, -2.2, 6.5, 6.5, fill=TEAL, shape=MSO_SHAPE.OVAL, alpha=18)
    box(s, -2.4, 4.6, 5.5, 5.5, fill=TEAL_L, shape=MSO_SHAPE.OVAL, alpha=10)
    text(s, 0.7, 1.0, 6, 0.3, "HACKATHON PITCH · 2026", size=12, color=TEAL_L, bold=True, spc=300)
    text(s, 0.7, 1.35, 6.2, 1.0, "Smart Pharmacy", size=56, bold=True, color=WHITE)
    text(s, 0.7, 2.45, 6.2, 0.7, "الصيدلية الذكية", size=32, bold=True, color=TEAL_L)
    text(s, 0.7, 3.35, 5.7, 1.2, "Find the nearest pharmacy that actually has your medicine — verified, fresh, in Arabic & English.", size=20, color="E2E8F0", line=1.1)
    for i, t in enumerate(["Patients", "Pharmacies", "Warehouses"]):
        b = box(s, 0.7 + i * 1.75, 4.75, 1.6, 0.46, line=TEAL_L, radius=0.5, lw=1.25)
        label(b, t, 13, WHITE, False)
    text(s, 0.7, 6.15, 6.2, 0.3, "Live prototype · demo data · Ismailia, Egypt", size=12, color="94A3B8")
    text(s, 0.7, 6.48, 6.6, 0.5, "Ahmed Alsaholy · Ahmed Shahen · Khaled Zakarya · Mohamed Ekramy · Fayez Mohamed", size=12, color="CBD5E1", bold=True)
    pic(s, "browser-map.png", 6.35, 0.85, w=6.7)
    pic(s, "phone-ar-results.png", 10.3, 2.3, h=5.05)
    notes(s, "Hi — we're team Smart Pharmacy. In Egypt, when a medicine goes missing, patients don't know which pharmacy still has it. "
             "We built a working platform that shows the nearest verified pharmacy with fresh stock, and connects pharmacies and warehouses so that stock data stays honest. "
             "Everything you'll see is a live prototype running on clearly labelled demo data.")

    # 2 — Problem
    s = new()
    header(s, "The problem", "When a medicine disappears, patients go pharmacy to pharmacy")
    box(s, 0.6, 1.7, 5.75, 3.4, fill=CORAL_L, line="FECDD3", radius=0.06)
    text(s, 0.95, 1.95, 5.1, 0.3, "ILLUSTRATIVE SCENARIO", size=10, color=CORAL, bold=True, spc=200)
    text(s, 0.95, 2.3, 5.15, 2.2, [
        "Um Ahmed, 58, lives in Ismailia and takes Concor 5 mg for her blood pressure.",
        "Her pharmacy is out. She calls six pharmacies, visits three — and still can't tell who has it, or when the next delivery arrives."], size=17, color=NAVY, line=1.1, after=10)
    text(s, 0.95, 4.45, 5.1, 0.5, "Concor, Augmentin, Brufen and Cataflam were all reported short in 2024 [5][11].", size=11.5, color=MUTED, italic=True)
    cards = [("Patients", CORAL, "Lost time and money; some turn to black-market packs of unknown origin [11]."),
             ("Pharmacists", AMBER, "About 30% of their time goes to checking stock and calling suppliers by hand [12]."),
             ("Warehouses", SKY, "No live view of where demand is — orders arrive by phone, shortages are spotted late.")]
    for i, (t, c, d) in enumerate(cards):
        y = 1.7 + i * 1.17
        box(s, 6.6, y, 6.15, 1.02, fill=WHITE, line=BORDER, radius=0.08, shadow=True)
        box(s, 6.6, y, 0.09, 1.02, fill=c, shape=MSO_SHAPE.RECTANGLE)
        text(s, 6.9, y + 0.14, 5.7, 0.3, t, size=15, bold=True, color=NAVY)
        text(s, 6.9, y + 0.47, 5.7, 0.5, d, size=12.5, color=SLATE)
    kpi(s, 0.6, 5.35, 3.95, 1.45, "~1,000", "of 17,000 registered medicines hit by shortages (Jul 2024) [5]", color=CORAL)
    kpi(s, 4.7, 5.35, 3.95, 1.45, "30M+", "Egyptians with hypertension depend on chronic medicines [5]", color=CORAL)
    kpi(s, 8.8, 5.35, 3.95, 1.45, "11M", "people live with diabetes and need insulin and tablets every month [11]", color=CORAL)
    footer(s, n, "[5] [11] [12]")
    notes(s, "This is a real pattern: in 2024 around a thousand medicines were reported short, including Concor, Augmentin, Brufen and Cataflam. "
             "Patients lose time and some buy from the black market. Pharmacists spend about 30% of their day chasing stock. Warehouses can't see demand.")

    # 3 — Shortages in numbers
    s = new()
    header(s, "The problem in numbers", "Shortages come and go — and nobody shares one live picture")
    text(s, 0.6, 1.6, 7, 0.3, "Reported number of medicines in shortage — Egypt", size=14, bold=True, color=NAVY)
    column_chart(s, 0.5, 1.95, 7.1, 4.1, ["Jul 2024", "Aug 2024", "Mar 2025*", "Sep 2025"], [1000, 500, 1471, 200], [CORAL, "FB7185", AMBER, "FDA4AF"], lsize=16, csize=12)
    text(s, 0.6, 6.12, 7.0, 0.6, "* Branded trade names (FEI Pharmaceutical Chamber). Counts come from different bodies and definitions — that inconsistency is part of the problem [5][6][8][9].", size=10, color=MUTED, italic=True)
    ins = [("20–40%", "of trade names were short in mid-2024, while 15–20% of active-ingredient demand went unmet [7]."),
           ("300", "branded drugs had no available alternative in March 2025 [8]."),
           ("40% → <5%", "shortage share by late 2025 — yet insulin, cardiac and oncology medicines still go missing [9][10].")]
    for i, (v, d) in enumerate(ins):
        y = 1.65 + i * 1.42
        box(s, 8.0, y, 4.75, 1.28, fill=LIGHT, line=BORDER, radius=0.08)
        text(s, 8.25, y + 0.12, 4.3, 0.5, v, size=24, bold=True, color=CORAL)
        text(s, 8.25, y + 0.6, 4.35, 0.65, d, size=11.5, color=SLATE)
    b = box(s, 8.0, 5.95, 4.75, 0.75, fill=NAVY, radius=0.1)
    label(b, "Patients need to know where the last boxes are — today.", 13, WHITE, True)
    footer(s, n, "[5]–[10]")
    notes(s, "Look at the numbers: depending on who counts, Egypt had between 200 and almost 1,500 medicines in shortage in the last two years. "
             "Even the official figures disagree — because nobody has a shared, live view of stock. Supply recovered by late 2025, but critical medicines still disappear.")

    # 4 — Root cause
    s = new()
    header(s, "Root cause", "95,000 pharmacies. Zero shared stock visibility.")
    grid = [("95,000", "community pharmacies — one of the highest densities in the world [2]", TEAL),
            ("≤ 2", "pharmacies per pharmacist by law → mostly independent shops, no chain-wide inventory [13]", TEAL),
            ("7,000+", "products from ~3,000 suppliers, still ordered largely by phone [12]", AMBER),
            ("~30%", "of a pharmacist's time spent checking stock and prices manually [12]", AMBER)]
    for i, (v, d, c) in enumerate(grid):
        kpi(s, 0.6 + (i % 2) * 3.15, 1.7 + (i // 2) * 2.35, 3.0, 2.2, v, d, color=c, vsize=34, dsize=12.5)
    box(s, 7.05, 1.7, 5.7, 4.55, fill=LIGHT, line=BORDER, radius=0.05)
    text(s, 7.35, 1.9, 5.2, 0.3, "Today: every link is a phone call", size=14, bold=True, color=NAVY)
    pt = box(s, 7.45, 3.55, 0.95, 0.95, fill=TEAL, shape=MSO_SHAPE.OVAL)
    label(pt, "Patient", 11)
    wh = box(s, 11.15, 3.5, 1.35, 1.05, fill=AMBER, radius=0.15)
    label(wh, "Warehouses", 11, NAVY)
    for i in range(5):
        y = 2.35 + i * 0.7
        ph = box(s, 9.35, y, 0.85, 0.5, fill=WHITE, line="94A3B8", radius=0.3)
        label(ph, "Rx ?", 11, MUTED)
        arrow(s, 8.42, 4.02, 9.33, y + 0.25, "94A3B8", 1.25, dash=True)
        arrow(s, 10.22, y + 0.25, 11.13, 4.02, "94A3B8", 1.25, dash=True)
    text(s, 7.35, 5.75, 5.2, 0.4, "No live stock  ·  no delivery ETAs  ·  no trust signal", size=12.5, color=CORAL, bold=True, align="c")
    footer(s, n, "[2] [12] [13]")
    notes(s, "Why does this happen? Egypt has around 95,000 pharmacies, mostly independent because the law limits ownership to two per pharmacist. "
             "There's no shared inventory system: stock lives in notebooks, isolated POS systems and phone calls to thousands of suppliers.")

    # 5 — Market & digital readiness
    s = new()
    header(s, "Opportunity", "A large, fast-growing and digital-ready market")
    text(s, 0.6, 1.6, 4.3, 0.3, "Egypt pharma market (EGP bn)", size=14, bold=True, color=NAVY)
    column_chart(s, 0.5, 1.95, 4.3, 3.9, ["2023", "2024", "2025"], [216, 309, 422], ["99F6E4", TEAL_L, TEAL], lsize=15, csize=12)
    text(s, 0.6, 5.9, 4.2, 0.5, "+43% in 2024 · +37% in 2025 — Egyptian Drug Authority [3][4]", size=10.5, color=MUTED, italic=True)
    text(s, 5.1, 1.6, 3.7, 0.3, "Registered pharmacists", size=14, bold=True, color=NAVY)
    column_chart(s, 5.0, 1.95, 3.7, 3.9, ["2003", "2013", "2023"], [71, 175, 313], ["BAE6FD", "38BDF8", SKY], fmt='0"k"', lsize=15, csize=12)
    text(s, 5.1, 5.9, 3.6, 0.5, "4.4× in 20 years — Pharmacists Syndicate records [2]", size=10.5, color=MUTED, italic=True)
    for i, (v, d) in enumerate([("109.5M", "people in Egypt (CAPMAS, Sep 2026) [1]"), ("96.3M", "internet users · 81.9% penetration (Jan 2025) [14]"),
                                ("92.6M", "mobile-internet subscriptions (Dec 2025) [15]"), ("~4B", "medicine packs sold in 2025 (EDA) [4]")]):
        kpi(s, 9.05, 1.6 + i * 1.22, 3.7, 1.1, v, d, vsize=24, dsize=11.5)
    footer(s, n, "[1]–[4] [14] [15]")
    notes(s, "The market is big and growing: EGP 422 billion in 2025 according to the Egyptian Drug Authority, with 313,000 pharmacists. "
             "And Egyptians are online: 96 million internet users, mostly on mobile — so a mobile-first, Arabic-first service can reach them.")

    # 6 — Solution
    s = new()
    header(s, "Our solution", "One verified network for medicine availability")
    text(s, 0.6, 1.42, 12, 0.5, "Smart Pharmacy connects patients, pharmacies and warehouses around one question: who has this exact medicine, right now, near me?", size=15, color=MUTED)
    cols = [("For patients", TEAL, ["Search by brand, generic, ingredient or barcode — Arabic or English", "Nearest verified pharmacies with fresh stock, 5 → 10 → 20 → 40 km", "Directions, call, ask to confirm or hold — or get an alert"]),
            ("For pharmacies", NAVY2, ["Inventory with batches, expiry and low-stock alerts", "Every unit change recorded in an auditable ledger", "Order from warehouses; track incoming stock and ETAs"]),
            ("For warehouses", "B45309", ["Accept or partially fill orders with automatic allocation", "Shipments, delivery dates and delays — with reasons", "Resolve discrepancies; see real demand signals"])]
    for i, (t, c, items) in enumerate(cols):
        x = 0.6 + i * 4.1
        box(s, x, 2.15, 3.95, 3.55, fill=WHITE, line=BORDER, radius=0.06, shadow=True)
        hdr = box(s, x, 2.15, 3.95, 0.62, fill=c, radius=0.2)
        label(hdr, t, 17, WHITE)
        bullet_list(s, x + 0.25, 3.0, 3.55, items, size=13, dot=c, gap=0.86)
    bar = box(s, 0.6, 5.95, 12.15, 0.72, fill=MINT, radius=0.2)
    label(bar, "Trust layer:  admin-verified pharmacies  ·  freshness labels  ·  auditable stock ledger  ·  privacy by design", 14.5, TEAL_D)
    footer(s, n)
    notes(s, "Our solution is one network for three users. Patients search and find verified stock nearby. Pharmacies manage inventory with a tamper-proof ledger and order from warehouses. "
             "Warehouses fulfil orders, schedule deliveries and see real demand. A trust layer makes sure what patients see is verified and fresh.")

    # 7 — How it works
    s = new()
    header(s, "How it works", "From search to pickup — and from warehouse to shelf")
    steps = [("1", "Search", "Exact product: brand, generic, ingredient or barcode — Arabic or English"),
             ("2", "Locate", "GPS (with consent), city / district, or a pin on the map"),
             ("3", "Rank", "Verified & fresh stock first; expand 5 → 10 → 20 → 40 km"),
             ("4", "Act", "Directions, call, confirm or hold — or set an alert")]
    for i, (num, t, d) in enumerate(steps):
        x = 0.6 + i * 3.1
        box(s, x, 1.6, 2.85, 1.55, fill=WHITE, line=BORDER, radius=0.08, shadow=True)
        c = box(s, x + 0.2, 1.78, 0.46, 0.46, fill=TEAL, shape=MSO_SHAPE.OVAL)
        label(c, num, 14)
        text(s, x + 0.78, 1.82, 1.9, 0.4, t, size=16, bold=True, color=NAVY)
        text(s, x + 0.2, 2.35, 2.5, 0.8, d, size=11.5, color=SLATE)
        if i < 3:
            arrow(s, x + 2.88, 2.37, x + 3.08, 2.37, TEAL, 2)
    box(s, 0.6, 3.45, 5.75, 3.25, fill=LIGHT, line=BORDER, radius=0.05)
    text(s, 0.85, 3.6, 5.3, 0.3, "Progressive radius search", size=14, bold=True, color=NAVY)
    cx, cy = 2.25, 5.25
    for r, lbl, a in [(1.35, "40 km", 8), (1.0, "20 km", 12), (0.68, "10 km", 18), (0.38, "5 km", 30)]:
        ring = box(s, cx - r, cy - r, 2 * r, 2 * r, fill=TEAL, line=TEAL, shape=MSO_SHAPE.OVAL, alpha=a, lw=0.75)
        text(s, cx + r * 0.72 - 0.25, cy - r * 0.72 - 0.12, 0.6, 0.22, lbl, size=8.5, color=TEAL_D, bold=True)
    box(s, cx - 0.09, cy - 0.09, 0.18, 0.18, fill=NAVY, shape=MSO_SHAPE.OVAL)
    for dx, dy, c in [(0.22, -0.12, TEAL), (-0.18, 0.2, TEAL), (0.55, 0.35, "94A3B8"), (-0.85, -0.4, "94A3B8"), (0.9, -0.75, "94A3B8"), (-0.4, 0.95, "94A3B8")]:
        box(s, cx + dx - 0.06, cy + dy - 0.06, 0.12, 0.12, fill=c, shape=MSO_SHAPE.OVAL)
    text(s, 3.85, 4.15, 2.35, 2.3, ["Stops at the first ring with verified, fresh stock.", "Patients can expand anytime — up to 40 km (configurable).", "One indexed query; results in milliseconds."], size=11.5, color=SLATE, after=6)
    box(s, 6.55, 3.45, 6.2, 3.25, fill=LIGHT, line=BORDER, radius=0.05)
    text(s, 6.8, 3.6, 5.8, 0.3, "Supply loop keeps stock honest", size=14, bold=True, color=NAVY)
    loop = [("Pharmacy orders", 0, 0), ("Warehouse confirms & allocates", 1, 0), ("Shipment + ETA (delays flagged)", 2, 0),
            ("Pharmacy receives & verifies", 2, 1), ("Ledger updates stock", 1, 1), ("Patients see it instantly", 0, 1)]
    pos = {}
    for t, col, row in loop:
        x, y = 6.8 + col * 1.98, 4.1 + row * 1.15
        b = box(s, x, y, 1.72, 0.78, fill=TEAL if t.startswith("Patients") else WHITE, line=TEAL, radius=0.15, lw=1.25)
        label(b, t, 11, WHITE if t.startswith("Patients") else NAVY, True)
        pos[t] = (x, y)
    arrow(s, 8.53, 4.49, 8.77, 4.49, TEAL, 1.75)
    arrow(s, 10.51, 4.49, 10.75, 4.49, TEAL, 1.75)
    arrow(s, 11.62, 4.9, 11.62, 5.23, TEAL, 1.75)
    arrow(s, 10.77, 5.64, 10.53, 5.64, TEAL, 1.75)
    arrow(s, 8.79, 5.64, 8.55, 5.64, TEAL, 1.75)
    text(s, 6.8, 6.17, 5.8, 0.4, "Orders and dispatches never inflate pharmacy stock — only a verified receipt does.", size=11, color=MUTED, italic=True)
    footer(s, n)
    notes(s, "A patient searches the exact product, shares a location or picks an area, and we rank verified, fresh stock first — starting at 5 km and expanding only if needed. "
             "Behind the scenes, the supply loop keeps data honest: orders and shipments never inflate stock; only a verified receipt updates the ledger, and patients see it immediately.")

    # 8 — Product: patients
    s = new(bg=LIGHT)
    header(s, "Product · Patients", "Patient finder — exact product, nearest verified stock")
    pic(s, "browser-finder-en.png", 0.35, 1.4, w=6.9)
    pic(s, "phone-ar-picker.png", 7.2, 1.3, h=5.3)
    pic(s, "phone-ar-results.png", 9.95, 1.3, h=5.3)
    text(s, 0.6, 6.55, 6.5, 0.45, "Distance · verified · open now · freshness — expected deliveries are kept separate, never counted as stock", size=11.5, color=SLATE, bold=True)
    text(s, 7.45, 6.55, 5.3, 0.45, "Exact-product picker (5 mg ≠ 2.5 mg) · Arabic-first, right-to-left, no account needed", size=11.5, color=SLATE, bold=True)
    footer(s, n)
    notes(s, "Here's the real app. The picker forces an exact choice — Concor 5 mg is never confused with 2.5 mg. "
             "Results show distance, a verified badge, opening status and how fresh the stock data is. Expected deliveries appear separately, so we never promise stock that isn't on the shelf.")

    # 9 — Product: pharmacies
    s = new(bg=LIGHT)
    header(s, "Product · Pharmacies", "Pharmacy workspace — inventory you can trust")
    pic(s, "card-inventory-table.png", 0.35, 1.35, w=6.55)
    pic(s, "card-incoming-table.png", 0.35, 3.45, w=6.55)
    pic(s, "card-ledger.png", 6.85, 1.3, w=6.1)
    f = box(s, 7.05, 4.55, 5.7, 0.8, fill=NAVY, radius=0.15)
    label(f, "Available = Physical − Reserved − Unavailable", 18, WHITE)
    text(s, 7.1, 5.42, 5.6, 0.5, "Never negative — enforced by the database. Every change is an append-only ledger entry.", size=11.5, color=SLATE)
    bullet_list(s, 0.6, 5.05, 6.2, ["Low-stock & expiry alerts · CSV import · barcode lookup", "Incoming stock tracked from order to verified receipt — separate from shelf stock", "Patient confirmation requests and holds with pickup codes"], size=12.5, gap=0.45)
    footer(s, n)
    notes(s, "For pharmacies: a clean inventory with clear availability, an incoming-stock view that tracks every order until verified receipt, and a ledger where every unit change is recorded. "
             "Available stock can never go negative — the database enforces it.")

    # 10 — Product: warehouses & admin
    s = new(bg=LIGHT)
    header(s, "Product · Warehouses & admin", "Warehouse operations and platform trust controls")
    pic(s, "card-warehouse-po.png", 0.35, 1.32, w=6.0)
    pic(s, "card-warehouse-shipment.png", 0.35, 4.0, w=6.0)
    pic(s, "browser-admin-overview.png", 6.45, 1.3, w=6.5)
    bullet_list(s, 6.75, 5.8, 6.0, ["Confirming an order reserves stock instantly — no overselling", "Delays require a reason; pharmacies are notified automatically", "Admins verify every pharmacy before patients can see it"], size=12, gap=0.38)
    footer(s, n)
    notes(s, "Warehouses confirm orders — fully or partially — and the stock is reserved immediately, so it can't be sold twice. Shipments carry delivery dates with a confidence level, and delays need a reason. "
             "Admins verify every pharmacy and see platform health, jobs and audit logs.")

    # 11 — Trust by design
    s = new()
    header(s, "Trust by design", "What patients see is verified, fresh — and provable")
    text(s, 0.6, 1.6, 5.0, 0.3, "Freshness windows (configurable)", size=14, bold=True, color=NAVY)
    rows = [("Update source", "Confirmed", "Outdated after"), ("Pharmacy POS integration", "2 h", "24 h"), ("Staff update · barcode · delivery receipt", "12 h", "72 h"), ("CSV import", "24 h", "7 days")]
    for i, (a, b2, c) in enumerate(rows):
        y = 2.0 + i * 0.58
        fill = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
        box(s, 0.6, y, 5.0, 0.54, fill=fill, line=BORDER if i else None, shape=MSO_SHAPE.RECTANGLE, lw=0.75)
        col = WHITE if i == 0 else SLATE
        text(s, 0.75, y, 2.8, 0.54, a, size=11.5, color=col, bold=i == 0, anchor="m")
        text(s, 3.6, y, 0.9, 0.54, b2, size=12, color=WHITE if i == 0 else TEAL, bold=True, anchor="m", align="c")
        text(s, 4.5, y, 1.05, 0.54, c, size=12, color=WHITE if i == 0 else CORAL, bold=True, anchor="m", align="c")
    text(s, 0.6, 4.45, 5.0, 0.9, "A failing or silent integration is automatically downgraded — it can never look fresh.", size=12, color=MUTED, italic=True)
    text(s, 5.95, 1.6, 3.2, 0.3, "What patients see", size=14, bold=True, color=NAVY)
    badges = [("Available", "DCFCE7", "166534"), ("Low stock", AMBER_L, "92400E"), ("Recently reported — unconfirmed", AMBER_L, "92400E"),
              ("Fewer units than requested", AMBER_L, "92400E"), ("Outdated — call first", CORAL_L, "9F1239"), ("Expected delivery — not in stock", "E0F2FE", "075985")]
    for i, (t, bg, fg) in enumerate(badges):
        b = box(s, 5.95, 2.0 + i * 0.55, 3.1, 0.42, fill=bg, radius=0.5)
        label(b, t, 11.5, fg)
    text(s, 9.45, 1.6, 3.3, 0.3, "Guarantees", size=14, bold=True, color=NAVY)
    bullet_list(s, 9.45, 2.0, 3.35, ["Available stock can't go negative (DB-enforced)", "Append-only ledger: every unit traceable", "Duplicate-safe orders & syncs", "Pharmacies can't touch each other's data", "Admin MFA, audit logs, rate limits", "Location used once, never stored"], size=11.5, gap=0.55)
    box(s, 0.6, 5.65, 12.15, 0.95, fill=MINT, radius=0.15)
    text(s, 0.9, 5.65, 11.6, 0.95, "Ranking rule: a closer but unconfirmed listing never outranks a farther verified and recently updated one.", size=15, color=TEAL_D, bold=True, anchor="m", align="c")
    footer(s, n, "[18]")
    notes(s, "Trust is the core. Every listing carries a freshness label based on its source: integrations are fresh for 2 hours, staff updates for 12. Outdated data is labelled and pushed down. "
             "The database guarantees stock can't go negative, every change is in an append-only ledger, and a closer unconfirmed listing never beats a verified one.")

    # 12 — Demo analytics
    s = new()
    header(s, "Data from the prototype", "Every search becomes a signal")
    out = A["outcomes"]
    found = out.get("FOUND", 0)
    total = sum(out.values())
    n_pharm = A.get("pharmacies", {}).get("n", 24)
    text(s, 0.6, 1.4, 12.2, 0.4, f"Simulated traffic: {total} patient searches across {n_pharm} demo pharmacies in Ismailia governorate — demo data, shown to illustrate the analytics, not market results.", size=12, color=MUTED, italic=True)
    text(s, 0.6, 1.95, 3.8, 0.3, "Search outcomes", size=14, bold=True, color=NAVY)
    doughnut(s, 0.45, 2.3, 3.9, 3.4, ["Found nearby", "No stock nearby"], [found, total - found], [TEAL, CORAL])
    b = A["bands"]
    text(s, 4.65, 1.95, 4.2, 0.3, "Where stock was found", size=14, bold=True, color=NAVY)
    column_chart(s, 4.55, 2.3, 4.25, 3.4, ["5 km", "10 km", "20 km", "40 km", "None"], [b.get("5", 0), b.get("10", 0), b.get("20", 0), b.get("40", 0), b.get("none", 0)], [TEAL, TEAL_L, "5EEAD4", "99F6E4", CORAL], lsize=13, csize=11, gap=45)
    fr = A["freshness"]
    text(s, 9.1, 1.95, 3.7, 0.3, "Freshness of listings", size=14, bold=True, color=NAVY)
    column_chart(s, 9.0, 2.3, 3.85, 3.4, ["Confirmed", "Unconfirmed", "Outdated", "POS-fed"], [fr.get("confirmed", 0), fr.get("recent", 0), fr.get("stale", 0), fr.get("integration", 0)], [TEAL, AMBER, CORAL, SKY], horizontal=True, lsize=13, csize=11, gap=45)
    unmet = list(A["unmet"].items())
    unmet_txt = f"all {unmet[0][1]} unmet searches were for {unmet[0][0]}" if len(unmet) == 1 else "unmet searches concentrate on: " + ", ".join(f"{k} ({v})" for k, v in unmet)
    box(s, 0.6, 5.9, 12.15, 0.8, fill=MINT, radius=0.15)
    text(s, 0.9, 5.9, 11.6, 0.8, [[("Insight: ", {"bold": True}), (f"{unmet_txt} — stocked only at a warehouse. That is the demand signal suppliers and authorities need, captured without storing personal data.", {})]], size=13, color=TEAL_D, anchor="m")
    footer(s, n, "[18] — live query of the prototype database")
    notes(s, "These charts come straight from our prototype's database. We simulated 400 patient searches on demo data. Most found stock within 5 km; outlying towns needed the wider rings. "
             "The interesting part: every unmet search was for insulin, which only a warehouse had. That's a demand signal — anonymous — that warehouses and authorities can act on.")

    # 13 — Proof
    s = new()
    header(s, "Proof it works", "A working prototype — measured, not mocked")
    tiles = [("67/67", "automated tests passing"), ("20/20", "end-to-end acceptance steps"), ("95", "documented REST API endpoints"),
             ("0", "stock-ledger mismatches"), ("0", "known dependency vulnerabilities"), ("AR · EN", "full right-to-left support")]
    for i, (v, d) in enumerate(tiles):
        kpi(s, 0.6 + (i % 3) * 2.2, 1.6 + (i // 3) * 1.6, 2.05, 1.45, v, d, vsize=28, dsize=11.5)
    text(s, 0.6, 4.95, 6.4, 0.3, "Automated tests by area", size=13, bold=True, color=NAVY)
    stacked_bar(s, 0.5, 5.2, 6.55, 1.55, [("Unit rules", 19), ("E2E workflow", 12), ("Search scenarios", 13), ("Supply & concurrency", 9), ("Security & jobs", 14)], [NAVY2, TEAL, TEAL_L, AMBER, SKY])
    text(s, 7.45, 1.6, 5.3, 0.3, "Finder response time (ms)", size=14, bold=True, color=NAVY)
    grouped_chart(s, 7.35, 1.95, 5.45, 3.85, ["Median (p50)", "95th percentile"], [("Sequential", [15, 20]), ("10 concurrent users", [84, 134])], [TEAL, AMBER])
    text(s, 7.45, 5.9, 5.3, 0.8, "Synthetic load: 3,000 pharmacies · 90,000 stock lines · 2-vCPU sandbox · 109 requests/s at 10 concurrent users.", size=11, color=MUTED, italic=True)
    footer(s, n, "[18]")
    notes(s, "This isn't a mockup. 67 automated tests pass, including a 20-step end-to-end acceptance test covering admin, pharmacy, warehouse and patient. "
             "Search answers in 15 milliseconds at the median on 3,000 synthetic pharmacies, and the stock ledger reconciles with zero mismatches.")

    # 14 — Competition
    s = new()
    header(s, "Competition", "Where we fit — and why we're different")
    mx, my, mw, mh = 0.75, 1.75, 6.3, 4.6
    box(s, mx, my, mw, mh, fill=LIGHT, line=BORDER, radius=0.03)
    box(s, mx + mw / 2 - 0.005, my, 0.01, mh, fill="CBD5E1", shape=MSO_SHAPE.RECTANGLE)
    box(s, mx, my + mh / 2 - 0.005, mw, 0.01, fill="CBD5E1", shape=MSO_SHAPE.RECTANGLE)
    text(s, mx, my + mh + 0.08, mw, 0.3, "Patient-facing   ←→   Pharmacy & supply-facing (B2B)", size=11, color=MUTED, bold=True, align="c")
    vt = text(s, mx - 2.55, my + mh / 2 - 0.15, 4.6, 0.3, "Ordering & delivery   ←→   Live availability & supply", size=11, color=MUTED, bold=True, align="c")
    vt.rotation = 270
    comps = [("Chefaa", 1.15, 4.75), ("Yodawy", 1.15, 5.35), ("Vezeeta", 2.35, 5.05), ("Pharmacy Marts", 4.35, 4.85), ("iSupply", 5.6, 5.45),
             ("Pharmacy chain apps", 1.05, 3.55), ("EDA 'Availability' reports", 1.05, 2.85)]
    for t, x, y in comps:
        b = box(s, x, y, 1.75 if len(t) > 10 else 1.2, 0.42, fill=WHITE, line="94A3B8", radius=0.5)
        label(b, t, 10.5, SLATE, False)
    star = box(s, 3.05, 2.05, 2.75, 0.62, fill=TEAL, radius=0.5, shadow=True)
    label(star, "Smart Pharmacy", 15, WHITE)
    text(s, 3.05, 2.72, 2.75, 0.5, "patients + pharmacies + warehouses, live", size=10, color=TEAL_D, bold=True, align="c")
    text(s, 7.45, 1.7, 5.3, 0.3, "What only we combine", size=15, bold=True, color=NAVY)
    diffs = ["Live, verified multi-pharmacy availability with freshness labels",
             "Progressive 5 → 40 km search; Arabic-first; works without an account or GPS",
             "Warehouse → pharmacy loop: incoming stock visible, receipts update stock",
             "Auditable stock ledger — no overselling, every unit traceable",
             "Anonymous demand signals (unmet searches) for suppliers & regulators"]
    for i, d in enumerate(diffs):
        y = 2.15 + i * 0.82
        c = box(s, 7.45, y, 0.42, 0.42, fill=TEAL, shape=MSO_SHAPE.OVAL)
        label(c, str(i + 1), 12)
        text(s, 8.05, y - 0.02, 4.7, 0.75, d, size=12.5, color=SLATE)
    text(s, 7.45, 6.3, 5.3, 0.45, "Positioning based on public company descriptions (2023–2026); competitors may offer features not shown — not a feature audit.", size=9.5, color=MUTED, italic=True)
    footer(s, n, "[12] [16] [17]")
    notes(s, "Existing players are strong at what they do: Chefaa, Yodawy and Vezeeta focus on ordering and delivery; Pharmacy Marts and iSupply digitise pharmacy procurement; "
             "the Drug Authority collects shortage reports. We sit across these: live, verified availability for patients, connected to the pharmacy and warehouse supply loop.")

    # 15 — Business model & GTM
    s = new()
    header(s, "Business model & go-to-market", "Free for patients. Valuable for the supply chain.")
    rev = [("Pharmacy Pro (SaaS)", "Free listing & stock updates; Pro adds ledger analytics, POS integration, team accounts and alerts.", TEAL),
           ("Warehouse plan", "Order & shipment management for distributors + a small fee per fulfilled order.", "B45309"),
           ("Demand insights", "Aggregated, anonymised shortage & demand signals for distributors, manufacturers and authorities.", SKY)]
    for i, (t, d, c) in enumerate(rev):
        y = 1.65 + i * 1.32
        box(s, 0.6, y, 6.3, 1.18, fill=WHITE, line=BORDER, radius=0.08, shadow=True)
        box(s, 0.6, y, 0.09, 1.18, fill=c, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.9, y + 0.13, 5.8, 0.3, t, size=15, bold=True, color=NAVY)
        text(s, 0.9, y + 0.47, 5.85, 0.7, d, size=12, color=SLATE)
    text(s, 0.6, 5.65, 6.3, 0.35, "Proposed model — to be validated during the pilot.", size=10.5, color=MUTED, italic=True)
    p = box(s, 0.6, 6.0, 6.3, 0.62, fill=NAVY, radius=0.2)
    label(p, "Patients always free · no selling of personal data", 13, WHITE)
    text(s, 7.35, 1.65, 5.4, 0.3, "Go-to-market", size=15, bold=True, color=NAVY)
    gtm = [("Pilot · Ismailia", "1.5M people [1] · 5–10 pharmacies + 1 warehouse · 8 weeks"), ("Suez Canal cities", "Port Said · Suez · Ismailia governorate towns"), ("Greater Cairo", "Partner with national distributors and pharmacy groups")]
    for i, (t, d) in enumerate(gtm):
        y = 2.1 + i * 1.12
        c = box(s, 7.35, y, 0.5, 0.5, fill=TEAL if i == 0 else WHITE, line=TEAL, shape=MSO_SHAPE.OVAL, lw=1.5)
        label(c, str(i + 1), 13, WHITE if i == 0 else TEAL)
        if i < 2:
            box(s, 7.59, y + 0.52, 0.02, 0.6, fill=TEAL_L, shape=MSO_SHAPE.RECTANGLE)
        text(s, 8.05, y + 0.02, 4.7, 0.3, t, size=14, bold=True, color=NAVY)
        text(s, 8.05, y + 0.35, 4.7, 0.6, d, size=11.5, color=SLATE)
    box(s, 7.35, 5.45, 5.4, 1.17, fill=MINT, radius=0.1)
    text(s, 7.6, 5.55, 5.0, 1.0, [[("Pilot KPIs  ", {"bold": True})], "Search success rate · time-to-medicine · % fresh listings · request response time · shortage signal lead time"], size=11.5, color=TEAL_D)
    footer(s, n, "[1]")
    notes(s, "Patients will always use it for free. Revenue comes from pharmacies upgrading to Pro features, warehouses using order and shipment management, and anonymised demand insights. "
             "We'll validate this in an 8-week pilot in Ismailia with 5 to 10 pharmacies and one warehouse, then expand along the Suez Canal and to Cairo.")

    # 16 — Roadmap
    s = new()
    header(s, "Roadmap", "Built in this hackathon — ready for a pilot")
    phases = [("Done · MVP", TEAL, ["Patient finder (AR / EN, RTL)", "Inventory ledger, CSV & barcode", "Orders, shipments & receiving", "Admin verification & audit logs", "67 automated tests passing"]),
              ("Next 3 months", AMBER, ["Ismailia pilot with real pharmacies", "SMS / WhatsApp availability alerts", "Password reset & e-mail verification", "POS connectors", "Verified catalog import (EDA data)"]),
              ("6–12 months", SKY, ["Demand forecasting & reorder suggestions", "Insurance & e-prescription partners", "Expansion to Greater Cairo", "Open API for health authorities", "Supplier performance scorecards"])]
    box(s, 0.9, 2.3, 11.6, 0.06, fill=BORDER, shape=MSO_SHAPE.RECTANGLE)
    for i, (t, c, items) in enumerate(phases):
        x = 0.6 + i * 4.1
        dot = box(s, x + 0.25, 2.1, 0.46, 0.46, fill=c, shape=MSO_SHAPE.OVAL, shadow=True)
        label(dot, "✓" if i == 0 else str(i + 1), 13)
        text(s, x + 0.85, 2.13, 3.0, 0.4, t, size=17, bold=True, color=NAVY)
        box(s, x, 2.85, 3.95, 3.6, fill=LIGHT, line=BORDER, radius=0.05)
        bullet_list(s, x + 0.25, 3.1, 3.55, items, size=12.5, dot=c, gap=0.62)
    footer(s, n)
    notes(s, "We've built the core this hackathon. Next: a real pilot in Ismailia, SMS and WhatsApp alerts, POS connectors and a verified catalog from Drug Authority data. "
             "Then forecasting, insurance partners and expansion to Cairo.")

    # 17 — Team
    s = new(dark=True)
    box(s, 9.0, -2.0, 6.0, 6.0, fill=TEAL, shape=MSO_SHAPE.OVAL, alpha=15)
    header(s, "Team", "The team behind Smart Pharmacy", dark=True)
    team = ["Ahmed Alsaholy", "Ahmed Shahen", "Khaled Zakarya", "Mohamed Ekramy", "Fayez Mohamed"]
    for i, name in enumerate(team):
        x = 0.85 + i * 2.45
        av = box(s, x + 0.35, 2.15, 1.45, 1.45, fill=TEAL if i % 2 == 0 else TEAL_D, line=TEAL_L, shape=MSO_SHAPE.OVAL, lw=2)
        label(av, "".join(w[0] for w in name.split()).upper(), 26)
        text(s, x - 0.2, 3.8, 2.55, 0.4, name, size=15, bold=True, color=WHITE, align="c")
    text(s, 0.6, 4.75, 12.1, 0.3, "BUILT WITH", size=11, color=TEAL_L, bold=True, spc=200, align="c")
    stack = ["NestJS", "PostgreSQL + PostGIS", "Next.js", "React", "TypeScript", "Tailwind CSS"]
    sx = (13.333 - (len(stack) * 1.85 - 0.15)) / 2
    for i, t in enumerate(stack):
        b = box(s, sx + i * 1.85, 5.1, 1.7, 0.45, line=TEAL_L, radius=0.5, lw=1)
        label(b, t, 11.5, WHITE, False)
    text(s, 0.6, 6.0, 12.1, 0.6, [[("Thank you — let's make every medicine findable.   ", {}), ("شكرًا", {"color": TEAL_L})]], size=20, color=WHITE, bold=True, align="c")
    footer(s, n, dark=True)
    notes(s, "We're a team of five: Ahmed Alsaholy, Ahmed Shahen, Khaled Zakarya, Mohamed Ekramy and Fayez Mohamed. Thank you — we'd love your questions, and we'd love to pilot this with real pharmacies.")

    # 18 — Sources
    s = new()
    header(s, "Sources", "Data sources")
    half = (len(SOURCES) + 1) // 2
    for col in range(2):
        paras = []
        for i, (t, url) in enumerate(SOURCES[col * half:(col + 1) * half], start=col * half + 1):
            paras.append([(f"[{i}]  ", {"bold": True, "color": TEAL}), (t, {})])
            paras.append([(url if url else "Internal: docs/10-testing-and-quality.md and live demo database", {"color": SKY if url else MUTED, "link": url, "size": 8})])
        text(s, 0.6 + col * 6.2, 1.5, 5.95, 5.3, paras, size=9.5, color=SLATE, line=1.0, after=2.5)
    text(s, 0.6, 6.75, 12.1, 0.3, "Market figures as reported by the cited sources (retrieved Oct 2026). Screens and demo analytics use clearly labelled demo data from the prototype; performance figures are team measurements on synthetic data.", size=9, color=MUTED, italic=True)
    footer(s, n)
    notes(s, "All market figures come from the sources listed here, with links. Prototype numbers are our own measurements, and every screen uses clearly labelled demo data.")

    prs.save(OUT)
    print(f"saved {OUT} ({OUT.stat().st_size // 1024} KB, {n} slides)")


if __name__ == "__main__":
    build()
